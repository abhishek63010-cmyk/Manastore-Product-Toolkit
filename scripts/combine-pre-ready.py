#!/usr/bin/env python3
"""
scripts/combine-pre-ready.py

Standalone consolidation script for MANA Store Product Toolkit.
Discovers catalogue-folder-named workbooks (or legacy 'manastore_pre_ready.xlsx' files)
in immediate catalogue subfolders
of a batch run directory, validates their schemas, and consolidates all product rows
into a single root-level '<BATCH_DIRECTORY>/manastore_pre_ready.xlsx'.

This script runs only when explicitly executed and NEVER modifies source workbooks.
"""

import sys
import os
import argparse
from pathlib import Path
from typing import List, Tuple, Dict, Any, Optional
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.cell.cell import Cell

# Add project root to sys.path so we can reuse schema definitions if available
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

try:
    from src.mappers.amra_transformer import AmraWorkbookTransformer
    EXPECTED_COLUMNS = list(AmraWorkbookTransformer.EXACT_COLUMNS)
except Exception:
    EXPECTED_COLUMNS = [
        "No",
        "supplier_name",
        "catalog_name",
        "product_group",
        "product_name",
        "varient_name",
        "supplier_sku",
        "product_type",
        "category_slug",
        "brand_name",
        "fabric",
        "blouse_fabric",
        "work",
        "occasion",
        "size",
        "description",
        "supplier_price",
        "gst_percent",
        "selling_price",
        "image-sku",
        "image_url",
        "publish",
        "image1_url",
        "image2_url",
        "image3_url",
        "image4_url",
    ]

FILENAME_TARGET = "manastore_pre_ready.xlsx"
LEGACY_FILENAME = FILENAME_TARGET
EXCLUDED_DIR_NAMES = {"_internal", "_reports"}


class ConsolidationError(Exception):
    """Custom exception for consolidation errors."""
    pass


def is_row_empty(row_values: Tuple[Any, ...]) -> bool:
    """Check if all cells in a row are None or empty/whitespace string."""
    for val in row_values:
        if val is not None:
            if isinstance(val, str) and val.strip() == "":
                continue
            return False
    return True


def discover_source_workbooks(batch_dir: Path) -> List[Path]:
    """
    Discover folder-named workbooks or legacy 'manastore_pre_ready.xlsx' files
    directly inside immediate catalogue subfolders of the batch directory.
    Excludes _internal, _reports, dot/underscore folders, and the batch root itself.
    """
    if not batch_dir.exists() or not batch_dir.is_dir():
        raise ConsolidationError(f"Batch directory does not exist or is not a directory: {batch_dir}")

    source_workbooks: List[Path] = []
    
    # Iterate immediate subdirectories only
    subdirs = sorted([d for d in batch_dir.iterdir() if d.is_dir()])
    
    for subdir in subdirs:
        # Ignore helper / internal / report directories
        if subdir.name in EXCLUDED_DIR_NAMES or subdir.name.startswith((".", "_")):
            continue
            
        candidate_names = (f"{subdir.name}.xlsx", LEGACY_FILENAME)
        candidate_file = next(
            (subdir / name for name in candidate_names if (subdir / name).is_file()),
            None,
        )
        if candidate_file is not None:
            source_workbooks.append(candidate_file)

    return source_workbooks


def validate_source_workbook(file_path: Path) -> Tuple[List[str], int]:
    """
    Validate a single source workbook:
    - Opens with openpyxl.
    - Has exact 26 Amra columns in exact order on the first row.
    - Has at least 1 non-empty product data row.
    - Returns (headers, product_row_count).
    """
    try:
        wb = openpyxl.load_workbook(file_path, data_only=False, read_only=False)
    except Exception as e:
        raise ConsolidationError(f"[{file_path.parent.name}] Failed to open workbook: {e}")

    try:
        ws = wb.active
        if ws is None:
            raise ConsolidationError(f"[{file_path.parent.name}] Workbook contains no active worksheet")

        # Read header row
        header_row = []
        for col_idx in range(1, len(EXPECTED_COLUMNS) + 1):
            cell_val = ws.cell(row=1, column=col_idx).value
            header_row.append(str(cell_val).strip() if cell_val is not None else "")

        if header_row != EXPECTED_COLUMNS:
            mismatches = []
            for i, (expected, actual) in enumerate(zip(EXPECTED_COLUMNS, header_row)):
                if expected != actual:
                    mismatches.append(f"Col {i+1}: expected '{expected}', found '{actual}'")
            if len(header_row) != len(EXPECTED_COLUMNS):
                mismatches.append(f"Column count mismatch: expected {len(EXPECTED_COLUMNS)}, found {len(header_row)}")
            
            error_details = "; ".join(mismatches[:5])
            if len(mismatches) > 5:
                error_details += f"; ... and {len(mismatches) - 5} more"
            raise ConsolidationError(
                f"[{file_path.parent.name}] Schema mismatch in '{file_path.name}'. {error_details}"
            )

        # Count non-empty data rows
        product_row_count = 0
        max_row = ws.max_row
        for r_idx in range(2, max_row + 1):
            row_vals = tuple(ws.cell(row=r_idx, column=c_idx).value for c_idx in range(1, len(EXPECTED_COLUMNS) + 1))
            if not is_row_empty(row_vals):
                product_row_count += 1

        if product_row_count == 0:
            raise ConsolidationError(
                f"[{file_path.parent.name}] Workbook '{file_path.name}' has 0 product data rows (empty workbook)"
            )

        return header_row, product_row_count
    finally:
        wb.close()


def consolidate_workbooks(
    batch_dir: Path,
    source_files: List[Path],
    output_path: Path,
    overwrite: bool = False,
    dry_run: bool = False
) -> Dict[str, Any]:
    """
    Validate and consolidate all source workbooks into output_path.
    """
    # Check output file safety
    if output_path.exists() and not overwrite and not dry_run:
        raise ConsolidationError(
            f"Destination workbook already exists at '{output_path}'. "
            f"Pass --force / --overwrite to overwrite, or specify a different --output."
        )

    # Pre-validation phase
    validation_results: List[Dict[str, Any]] = []
    total_expected_rows = 0

    print("================================================================================")
    print("MANA STORE PRE-READY WORKBOOK CONSOLIDATION")
    print("================================================================================")
    print(f"Batch Directory: {batch_dir.resolve()}")
    print(f"Discovered {len(source_files)} catalogue workbook(s):")
    print("--------------------------------------------------------------------------------")

    for idx, sfile in enumerate(source_files, 1):
        headers, row_count = validate_source_workbook(sfile)
        validation_results.append({
            "index": idx,
            "catalogue": sfile.parent.name,
            "path": sfile,
            "rows": row_count
        })
        total_expected_rows += row_count
        print(f"  [{idx:02d}] {sfile.parent.name:<30} -> {row_count:>4} row(s)  ({sfile.name})")

    print("--------------------------------------------------------------------------------")
    print(f"Total Discovered Product Rows: {total_expected_rows}")
    print(f"Target Output Path:            {output_path.resolve()}")
    print("================================================================================")

    if dry_run:
        print("[DRY-RUN] Validation passed for all workbooks. No output file was written.")
        return {
            "status": "DRY_RUN",
            "source_count": len(source_files),
            "catalogues": validation_results,
            "total_rows": total_expected_rows,
            "output_path": str(output_path.resolve())
        }

    # Consolidation phase using openpyxl
    combined_wb = openpyxl.Workbook()
    combined_ws = combined_wb.active
    combined_ws.title = "Products"

    # Write header once (using style and width hints from the first workbook if possible)
    first_file = source_files[0]
    first_wb = openpyxl.load_workbook(first_file, data_only=False)
    first_ws = first_wb.active

    # Copy header values and basic styling from first workbook
    for c_idx in range(1, len(EXPECTED_COLUMNS) + 1):
        src_cell = first_ws.cell(row=1, column=c_idx)
        dest_cell = combined_ws.cell(row=1, column=c_idx, value=src_cell.value)
        if src_cell.has_style:
            dest_cell.font = openpyxl.styles.Font(
                name=src_cell.font.name,
                size=src_cell.font.size,
                bold=src_cell.font.bold,
                color=src_cell.font.color
            )
            dest_cell.alignment = openpyxl.styles.Alignment(
                horizontal=src_cell.alignment.horizontal,
                vertical=src_cell.alignment.vertical
            )
        col_letter = openpyxl.utils.get_column_letter(c_idx)
        if col_letter in first_ws.column_dimensions:
            combined_ws.column_dimensions[col_letter].width = first_ws.column_dimensions[col_letter].width

    first_wb.close()

    # Append data rows
    current_dest_row = 2
    for sfile in source_files:
        wb = openpyxl.load_workbook(sfile, data_only=False)
        ws = wb.active
        for r_idx in range(2, ws.max_row + 1):
            row_vals = tuple(ws.cell(row=r_idx, column=c_idx).value for c_idx in range(1, len(EXPECTED_COLUMNS) + 1))
            if is_row_empty(row_vals):
                continue

            for c_idx in range(1, len(EXPECTED_COLUMNS) + 1):
                src_cell = ws.cell(row=r_idx, column=c_idx)
                dest_cell = combined_ws.cell(row=current_dest_row, column=c_idx, value=src_cell.value)
                if src_cell.number_format:
                    dest_cell.number_format = src_cell.number_format
            current_dest_row += 1
        wb.close()

    total_written_rows = current_dest_row - 2
    if total_written_rows != total_expected_rows:
        raise ConsolidationError(
            f"Row count mismatch during consolidation: expected {total_expected_rows}, wrote {total_written_rows}"
        )

    # Save to destination
    output_path.parent.mkdir(parents=True, exist_ok=True)
    combined_wb.save(output_path)
    combined_wb.close()

    # Generate consolidation_summary.txt beside output workbook
    summary_txt_path = output_path.parent / "consolidation_summary.txt"
    summary_lines = [
        "================================================================================",
        "MANA STORE PRE-READY CONSOLIDATION SUMMARY",
        "================================================================================",
        f"Consolidation Status: SUCCESS",
        f"Output Workbook:      {output_path.resolve()}",
        f"Source Workbooks:     {len(source_files)}",
        f"Total Product Rows:   {total_written_rows}",
        "--------------------------------------------------------------------------------",
        "SOURCE WORKBOOK BREAKDOWN:",
    ]
    for cat in validation_results:
        summary_lines.append(f"  - [{cat['index']:02d}] {cat['catalogue']:<30} : {cat['rows']} row(s)")
    summary_lines.append("================================================================================\n")

    with open(summary_txt_path, "w", encoding="utf-8") as f:
        f.write("\n".join(summary_lines))

    print(f"[SUCCESS] Combined {total_written_rows} rows from {len(source_files)} catalogues into:")
    print(f"          -> {output_path.resolve()}")
    print(f"          -> Summary saved: {summary_txt_path.resolve()}")

    return {
        "status": "SUCCESS",
        "source_count": len(source_files),
        "catalogues": validation_results,
        "total_rows": total_written_rows,
        "output_path": str(output_path.resolve()),
        "summary_path": str(summary_txt_path.resolve())
    }


def main():
    parser = argparse.ArgumentParser(
        description="Consolidate catalogue-folder-named or legacy pre-ready workbooks into one root workbook."
    )
    parser.add_argument(
        "batch_dir",
        type=str,
        help="Path to the batch run output directory containing catalogue subfolders."
    )
    parser.add_argument(
        "--output",
        "-o",
        type=str,
        default=None,
        help="Optional custom output path for the combined workbook (default: <batch_dir>/manastore_pre_ready.xlsx)."
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Validate source workbooks and preview row counts without writing the combined workbook."
    )
    parser.add_argument(
        "--force",
        "--overwrite",
        action="store_true",
        dest="overwrite",
        help="Overwrite the existing combined output workbook if present."
    )

    args = parser.parse_args()

    batch_dir = Path(args.batch_dir).resolve()
    if not batch_dir.exists():
        print(f"ERROR: Batch directory not found: {batch_dir}", file=sys.stderr)
        sys.exit(1)

    if args.output:
        output_path = Path(args.output).resolve()
    else:
        output_path = batch_dir / FILENAME_TARGET

    try:
        source_files = discover_source_workbooks(batch_dir)
        if not source_files:
            raise ConsolidationError(
                f"No folder-named or legacy '{LEGACY_FILENAME}' catalogue workbooks found in immediate subfolders of '{batch_dir}'. "
                f"(Subfolders checked must be direct subdirectories and not '_internal' or '_reports')"
            )

        consolidate_workbooks(
            batch_dir=batch_dir,
            source_files=source_files,
            output_path=output_path,
            overwrite=args.overwrite,
            dry_run=args.dry_run
        )
        sys.exit(0)

    except ConsolidationError as ce:
        print(f"\n[ERROR] Consolidation aborted: {ce}", file=sys.stderr)
        sys.exit(1)
    except Exception as e:
        print(f"\n[UNEXPECTED ERROR] {e}", file=sys.stderr)
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    main()
