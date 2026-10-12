import openpyxl
from pathlib import Path
import pytest

from src.pipeline.config import PipelineConfig
from src.pipeline.orchestrator import PipelineOrchestrator
from src.pipeline.output_manager import OutputManager


def test_simplified_output_folder_structure(tmp_path: Path):
    """
    Verify the simplified output folder structure:
    1. amra_data_ready.xlsx is the single admin-import file at the run-folder root.
    2. RUN_SUMMARY.txt is at the run-folder root with explicit admin instructions.
    3. _internal/ contains canonical models, raw extracted TSV, refined workbook, and manifest.
    4. _reports/ contains validation and audit reports.
    5. No duplicate workbooks or debug files alongside amra_data_ready.xlsx.
    """
    # Create sample apparel CSV input
    csv_file = tmp_path / "sample_catalog.csv"
    csv_file.write_text(
        "Catalogue Name,Product Name,Brand,Fabric,Work,Supplier Price,Total Designs\n"
        "Meera Vol 1,Meera Designer Saree,Meera,Silk,Embroidered,750,2 Pcs\n",
        encoding="utf-8",
    )

    run_output_dir = tmp_path / "run_folder"

    config = PipelineConfig(
        input_path=csv_file,
        output_dir=run_output_dir,
        supplier="generic",
        export_formats=["xlsx", "json", "csv"],
        generate_report=True,
    )

    result = PipelineOrchestrator.run(config)
    assert result.success is True

    # 1. Root folder verification
    root_items = [p.name for p in run_output_dir.iterdir()]
    # Root should ONLY have: amra_data_ready.xlsx, RUN_SUMMARY.txt, _internal, _reports
    assert "amra_data_ready.xlsx" in root_items
    assert "RUN_SUMMARY.txt" in root_items
    assert "_internal" in root_items
    assert "_reports" in root_items
    assert len(root_items) == 4, f"Expected exactly 4 items in root, got: {root_items}"

    # Verify amra_data_ready.xlsx has valid schema
    amra_wb = openpyxl.load_workbook(run_output_dir / "amra_data_ready.xlsx")
    assert len(amra_wb.sheetnames) >= 1
    ws = amra_wb.active
    headers = [cell.value for cell in ws[1]]
    assert len(headers) == 26
    assert headers[0] == "No"
    assert headers[18] == "selling_price"

    # Verify RUN_SUMMARY.txt content
    summary_text = (run_output_dir / "RUN_SUMMARY.txt").read_text(encoding="utf-8")
    assert "FINAL ADMIN IMPORT FILE: amra_data_ready.xlsx" in summary_text
    assert "_internal/" in summary_text
    assert "_reports/" in summary_text

    # 2. _internal/ directory verification
    internal_dir = run_output_dir / "_internal"
    assert internal_dir.is_dir()
    internal_files = [p.name for p in internal_dir.iterdir()]
    assert "canonical_products.xlsx" in internal_files
    assert "canonical_products.json" in internal_files
    assert "canonical_products.csv" in internal_files
    assert "refined_products.xlsx" in internal_files
    assert "mana_store_import_ready.xlsx" in internal_files
    assert "raw_supplier_extracted.tsv" in internal_files
    # Manifest exists if approval_manifest was generated
    if result.approval_manifest:
        assert "image_approval_manifest.json" in internal_files

    # Verify raw TSV contents
    raw_tsv_text = (internal_dir / "raw_supplier_extracted.tsv").read_text(encoding="utf-8")
    assert "Meera Vol 1" in raw_tsv_text

    # 3. _reports/ directory verification
    reports_dir = run_output_dir / "_reports"
    assert reports_dir.is_dir()
    reports_files = [p.name for p in reports_dir.iterdir()]
    assert "validation_report.xlsx" in reports_files
    assert "validation_report.json" in reports_files
