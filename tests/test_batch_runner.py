import csv
import json
import shutil
from pathlib import Path
import subprocess
import sys
import openpyxl
import pytest


PROJECT_ROOT = Path(__file__).resolve().parents[1]
RUN_BATCH_SCRIPT = PROJECT_ROOT / "scripts" / "run-batch.ps1"


def run_batch_powershell(manifest_path: Path, validate_only: bool = False) -> subprocess.CompletedProcess:
    """Helper to execute scripts/run-batch.ps1 via PowerShell."""
    cmd = [
        "powershell",
        "-ExecutionPolicy", "Bypass",
        "-File", str(RUN_BATCH_SCRIPT),
        "-ManifestPath", str(manifest_path),
    ]
    if validate_only:
        cmd.append("-ValidateOnly")

    return subprocess.run(cmd, cwd=str(PROJECT_ROOT), capture_output=True, text=True)


def test_batch_runner_validate_only_mode(tmp_path: Path):
    """Verify -ValidateOnly validates valid manifest without executing main.py."""
    # Create fake sample catalogue
    sample_csv = tmp_path / "valid_sample.csv"
    sample_csv.write_text("Product Name,SKU,Price\nTest Kurti,TK-01,500\n", encoding="utf-8")

    manifest_data = {
        "batchName": "Test_Validate_Batch",
        "catalogues": [
            {
                "name": "Cat1_Valid",
                "supplierFolder": "SupplierA",
                "supplierAdapter": "generic",
                "source": str(sample_csv),
                "images": None,
            }
        ]
    }
    manifest_file = tmp_path / "manifest.json"
    manifest_file.write_text(json.dumps(manifest_data), encoding="utf-8")

    proc = run_batch_powershell(manifest_file, validate_only=True)
    assert proc.returncode == 0
    assert "Manifest validation: PASSED" in proc.stdout
    assert "ValidateOnly switch was requested. Skipping execution." in proc.stdout


def test_batch_runner_rejects_empty_manifest(tmp_path: Path):
    """Verify empty catalogue list is rejected when no supplier roots are discoverable."""
    manifest_file = tmp_path / "empty_manifest.json"
    manifest_file.write_text(json.dumps({"batchName": "Empty", "catalogues": []}), encoding="utf-8")

    input_root = PROJECT_ROOT / "input"
    amra_root = input_root / "Amra_wholesale"
    mudhra_root = input_root / "Mudhra_Ethnic"
    renamed_roots = []

    for candidate in (amra_root, mudhra_root):
        if candidate.exists():
            backup = candidate.with_name(candidate.name + "_hidden_for_empty_manifest_test")
            shutil.move(str(candidate), str(backup))
            renamed_roots.append((candidate, backup))

    try:
        proc = run_batch_powershell(manifest_file, validate_only=True)
        assert proc.returncode != 0
        assert "at least one catalogue entry" in proc.stdout or "VALIDATION FAILED" in proc.stdout
    finally:
        for original, backup in renamed_roots:
            if backup.exists():
                shutil.move(str(backup), str(original))


def test_batch_runner_rejects_duplicate_catalogue_names(tmp_path: Path):
    """Verify duplicate catalogue names in manifest are detected and rejected."""
    sample_csv = tmp_path / "sample.csv"
    sample_csv.write_text("Name,SKU\nP,S\n", encoding="utf-8")

    manifest_data = {
        "batchName": "Duplicate_Test",
        "catalogues": [
            {
                "name": "SameName",
                "supplierAdapter": "generic",
                "source": str(sample_csv),
            },
            {
                "name": "SameName",
                "supplierAdapter": "generic",
                "source": str(sample_csv),
            }
        ]
    }
    manifest_file = tmp_path / "duplicate_manifest.json"
    manifest_file.write_text(json.dumps(manifest_data), encoding="utf-8")

    proc = run_batch_powershell(manifest_file, validate_only=True)
    assert proc.returncode != 0
    assert "Duplicate catalogue name detected: 'SameName'" in proc.stdout


def test_batch_runner_rejects_invalid_source_path(tmp_path: Path):
    """Verify non-existent source file is rejected with clear path error."""
    manifest_data = {
        "batchName": "Missing_Source_Test",
        "catalogues": [
            {
                "name": "MissingSourceCat",
                "supplierAdapter": "generic",
                "source": str(tmp_path / "non_existent_file.csv"),
            }
        ]
    }
    manifest_file = tmp_path / "missing_source.json"
    manifest_file.write_text(json.dumps(manifest_data), encoding="utf-8")

    proc = run_batch_powershell(manifest_file, validate_only=True)
    assert proc.returncode != 0
    assert "Source file does not exist" in proc.stdout


def test_batch_runner_rejects_unsupported_adapter(tmp_path: Path):
    """Verify unsupported supplier adapter is rejected."""
    sample_csv = tmp_path / "sample.csv"
    sample_csv.write_text("Name,SKU\nP,S\n", encoding="utf-8")

    manifest_data = {
        "batchName": "Bad_Adapter_Test",
        "catalogues": [
            {
                "name": "BadAdapterCat",
                "supplierAdapter": "unsupported_xyz_adapter",
                "source": str(sample_csv),
            }
        ]
    }
    manifest_file = tmp_path / "bad_adapter.json"
    manifest_file.write_text(json.dumps(manifest_data), encoding="utf-8")

    proc = run_batch_powershell(manifest_file, validate_only=True)
    assert proc.returncode != 0
    assert "Unsupported supplierAdapter 'unsupported_xyz_adapter'" in proc.stdout


def test_batch_runner_rejects_unresolved_placeholder(tmp_path: Path):
    """Verify unresolved placeholder <SupplierName> is rejected."""
    sample_csv = tmp_path / "sample.csv"
    sample_csv.write_text("Name,SKU\nP,S\n", encoding="utf-8")

    manifest_data = {
        "batchName": "Placeholder_Test",
        "catalogues": [
            {
                "name": "<SupplierB_Catalogue01>",
                "supplierAdapter": "generic",
                "source": str(sample_csv),
            }
        ]
    }
    manifest_file = tmp_path / "placeholder_manifest.json"
    manifest_file.write_text(json.dumps(manifest_data), encoding="utf-8")

    proc = run_batch_powershell(manifest_file, validate_only=True)
    assert proc.returncode != 0
    assert "unresolved placeholder" in proc.stdout


def test_batch_runner_execution_multiple_catalogues(tmp_path: Path):
    """
    Integration test: Run batch with multiple catalogues.
    Verify:
    1. Output directory is uniquely named by batchName.
    2. Subfolders for each catalogue created.
    3. Final workbook uses the catalogue output-folder name.
    4. batch_summary.csv is produced.
    """
    # Catalogue 1: generic apparel CSV
    csv1 = tmp_path / "cat1.csv"
    csv1.write_text(
        "Catalogue Name,Product Name,Brand,Supplier Price,Total Designs\n"
        "Aura Summer,Aura Straight Kurti,WeaveStyle,420,2 Pcs\n",
        encoding="utf-8",
    )

    # Catalogue 2: amra tabular format
    csv2 = tmp_path / "cat2.csv"
    csv2.write_text(
        "Catalogue Name,Product Name,Brand,Supplier Price,Total Designs\n"
        "Raas Leela,Raas Cotton Saree,Raas,650,2 Pcs\n",
        encoding="utf-8",
    )

    batch_name = "Multi_Cat_Test_Run"
    manifest_data = {
        "batchName": batch_name,
        "catalogues": [
            {
                "name": "Aura_Collection",
                "supplierFolder": "WeaveStyle",
                "supplierAdapter": "generic",
                "source": str(csv1),
                "images": None,
            },
            {
                "name": "Raas_Collection",
                "supplierFolder": "Raas",
                "supplierAdapter": "generic",
                "source": str(csv2),
                "images": None,
            }
        ]
    }
    manifest_file = tmp_path / "multi_manifest.json"
    manifest_file.write_text(json.dumps(manifest_data), encoding="utf-8")

    proc = run_batch_powershell(manifest_file, validate_only=False)
    assert proc.returncode == 0
    assert "All configured catalogues processed successfully!" in proc.stdout

    # Find created batch output directory
    output_root = PROJECT_ROOT / "output"
    matching_dirs = [d for d in output_root.iterdir() if d.is_dir() and d.name.startswith(batch_name)]
    assert len(matching_dirs) >= 1
    latest_batch_dir = sorted(matching_dirs, key=lambda d: d.name)[-1]

    # Check batch_summary.csv
    summary_csv = latest_batch_dir / "batch_summary.csv"
    assert summary_csv.exists()
    with open(summary_csv, encoding="utf-8") as f:
        rows = list(csv.reader(f))
    # Header + 2 catalogue rows
    assert len(rows) >= 3
    assert rows[1][1] == "Aura_Collection"
    assert rows[1][2] == "SUCCESS"
    assert rows[2][1] == "Raas_Collection"
    assert rows[2][2] == "SUCCESS"

    # Check catalogue 1 subfolder
    cat1_dir = latest_batch_dir / "01_Aura_Collection"
    cat1_workbook = cat1_dir / "01_Aura_Collection.xlsx"
    assert cat1_workbook.exists()
    assert (cat1_dir / "RUN_SUMMARY.txt").exists()
    assert (cat1_dir / "_internal").exists()
    assert (cat1_dir / "_reports").exists()
    # Ensure amra_data_ready.xlsx was renamed
    assert not (cat1_dir / "amra_data_ready.xlsx").exists()
    summary_text = (cat1_dir / "RUN_SUMMARY.txt").read_text(encoding="utf-8")
    assert "01_Aura_Collection.xlsx" in summary_text
    assert "manastore_pre_ready.xlsx" not in summary_text
    assert rows[1][9].endswith("01_Aura_Collection.xlsx")
    final_wb = openpyxl.load_workbook(cat1_workbook, read_only=True)
    assert final_wb.sheetnames == ["Amra_Data"]
    assert [cell.value for cell in next(final_wb.active.iter_rows(min_row=1, max_row=1))] == [
        "No",
        "supplier_name",
        "catalog_name",
        "product_group",
        "product_name",
        "varient_name",
        "supplier_sku",
        "product_type",
        "category_slug",
        "brand_name",
        "fabric",
        "blouse_fabric",
        "work",
        "occasion",
        "size",
        "description",
        "supplier_price",
        "gst_percent",
        "selling_price",
        "image-sku",
        "image_url",
        "publish",
        "image1_url",
        "image2_url",
        "image3_url",
        "image4_url",
    ]
    final_wb.close()

    # Check catalogue 2 subfolder
    cat2_dir = latest_batch_dir / "02_Raas_Collection"
    assert (cat2_dir / "02_Raas_Collection.xlsx").exists()
    assert (cat2_dir / "RUN_SUMMARY.txt").exists()
    assert (cat2_dir / "_internal").exists()
    assert (cat2_dir / "_reports").exists()


def test_batch_runner_failed_catalogue_followed_by_success(tmp_path: Path):
    """
    Verify runner continues to next catalogue when one catalogue fails,
    records failure in batch_summary.csv, and exits with nonzero code.
    """
    # 1. Broken catalogue (invalid CSV structure that fails pipeline)
    broken_csv = tmp_path / "broken.csv"
    broken_csv.write_text("", encoding="utf-8")  # Empty file fails extraction

    # 2. Good catalogue
    good_csv = tmp_path / "good.csv"
    good_csv.write_text(
        "Catalogue Name,Product Name,Brand,Supplier Price,Total Designs\n"
        "Shubh Silk,Shubh Festive Saree,Shubh,950,2 Pcs\n",
        encoding="utf-8",
    )

    batch_name = "Failure_Recovery_Batch"
    manifest_data = {
        "batchName": batch_name,
        "catalogues": [
            {
                "name": "Broken_Catalogue",
                "supplierAdapter": "generic",
                "source": str(broken_csv),
            },
            {
                "name": "Good_Catalogue",
                "supplierAdapter": "generic",
                "source": str(good_csv),
            }
        ]
    }
    manifest_file = tmp_path / "mixed_manifest.json"
    manifest_file.write_text(json.dumps(manifest_data), encoding="utf-8")

    proc = run_batch_powershell(manifest_file, validate_only=False)
    # Overall exit code must be non-zero because catalogue #1 failed
    assert proc.returncode != 0
    assert "One or more catalogues failed during batch execution." in proc.stdout

    # Find created batch output directory
    output_root = PROJECT_ROOT / "output"
    matching_dirs = [d for d in output_root.iterdir() if d.is_dir() and d.name.startswith(batch_name)]
    assert len(matching_dirs) >= 1
    latest_batch_dir = sorted(matching_dirs, key=lambda d: d.name)[-1]

    # Verify summary CSV reflects 1 FAILED and 1 SUCCESS
    summary_csv = latest_batch_dir / "batch_summary.csv"
    assert summary_csv.exists()
    with open(summary_csv, encoding="utf-8") as f:
        rows = list(csv.reader(f))

    assert len(rows) >= 3
    assert rows[1][1] == "Broken_Catalogue"
    assert rows[1][2] == "FAILED"
    assert rows[2][1] == "Good_Catalogue"
    assert rows[2][2] == "SUCCESS"

    # Good catalogue produced a folder-named workbook.
    assert (latest_batch_dir / "02_Good_Catalogue" / "02_Good_Catalogue.xlsx").exists()
